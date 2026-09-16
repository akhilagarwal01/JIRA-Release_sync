#!/usr/bin/env python3
"""
Pull COR/LINUX/ALPINE board tickets for configured status transitions and DEVOPS release tickets,
then append rows to the local DailyTaskLogs.xlsx workbook (quarterly sheet tab).

Weekdays only (Mon–Fri). Lookback: Monday = 3 days, Tue–Fri = 1 day.
"""

from __future__ import annotations

import sys
from pathlib import Path

_VENDOR = Path(__file__).resolve().parent / ".vendor"
if _VENDOR.is_dir():
    sys.path.insert(0, str(_VENDOR))

import argparse
import os
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from typing import Any

import requests
from dotenv import load_dotenv

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.worksheet import Worksheet
except ModuleNotFoundError:
    print(
        "Missing openpyxl. Install dependencies with:\n"
        "  python3 -m pip install -r requirements.txt\n"
        "Or use the wrapper:\n"
        "  bash scripts/run_daily_task_log.sh",
        file=sys.stderr,
    )
    sys.exit(1)

from sync import (
    _require_env,
    apply_jql_lookback_days,
    dedupe_issues_by_key,
    jira_get_issue_comments,
    jira_search_issues,
    jql_updated_lookback_start,
    parse_env_csv,
    release_issue_passes_creator_link_filter,
    saathi_release_calendar_days_newest_first,
    skip_automation_only_in_updated_window,
)

SHEET_HEADERS = ("Date", "TaskId", "Description", "Remarks", "QA Check")
SHEET_COLUMN_COUNT = len(SHEET_HEADERS)

DEFAULT_WORKBOOK = Path(__file__).resolve().parent / "DailyTaskLogs.xlsx"
DEFAULT_QA_MISSING_NOTE = "QA done by Peer or\nNo QA done"
DEFAULT_BOARD_TO_STATUSES = ("UAT/Staging",)
DEFAULT_BOARD_FROM_STATUSES = ("In QA",)
DEFAULT_BOARD_PROJECTS = ("COR", "LINUX", "ALPINE")


def board_projects() -> tuple[str, ...]:
    """Comma-separated project keys from DAILY_TASK_LOG_BOARD_PROJECTS (default: COR,LINUX,ALPINE)."""
    configured = parse_env_csv("DAILY_TASK_LOG_BOARD_PROJECTS")
    if configured:
        return tuple(project.upper() for project in configured)
    return DEFAULT_BOARD_PROJECTS


def board_to_statuses() -> tuple[str, ...]:
    """Comma-separated destination statuses from DAILY_TASK_LOG_BOARD_TO_STATUSES."""
    configured = parse_env_csv("DAILY_TASK_LOG_BOARD_TO_STATUSES")
    if configured:
        return tuple(configured)
    return DEFAULT_BOARD_TO_STATUSES


def board_from_statuses(to_statuses: tuple[str, ...]) -> tuple[str, ...] | None:
    """
    Optional comma-separated source statuses from DAILY_TASK_LOG_BOARD_FROM_STATUSES.
    Unset + UAT/Staging only → In QA (legacy). Unset + multiple to-statuses → any from.
    Explicit empty value → any from-status.
    """
    if "DAILY_TASK_LOG_BOARD_FROM_STATUSES" in os.environ:
        configured = parse_env_csv("DAILY_TASK_LOG_BOARD_FROM_STATUSES")
        return tuple(configured) if configured else None
    if to_statuses == DEFAULT_BOARD_TO_STATUSES:
        return DEFAULT_BOARD_FROM_STATUSES
    return None


def describe_board_transitions(
    to_statuses: tuple[str, ...],
    from_statuses: tuple[str, ...] | None,
) -> str:
    if from_statuses:
        pairs = [f"{src} → {dst}" for src in from_statuses for dst in to_statuses if src != dst]
        return ", ".join(pairs) if pairs else ", ".join(to_statuses)
    return ", ".join(f"→ {status}" for status in to_statuses)


@dataclass(frozen=True)
class TaskEntry:
    event_date: date
    task_id: str
    description: str
    remarks: str
    qa_check: str
    sort_rank: int  # COR/LINUX/ALPINE before DEVOPS within the same day


def qa_missing_note() -> str:
    raw = os.environ.get("DAILY_TASK_LOG_QA_MISSING_NOTE", DEFAULT_QA_MISSING_NOTE).strip()
    return raw.replace("\\n", "\n") if raw else DEFAULT_QA_MISSING_NOTE


def qa_to_statuses() -> frozenset[str]:
    raw = os.environ.get("DAILY_TASK_LOG_QA_TO_STATUSES", "Deployed on-FT,QA Signed OFF")
    return frozenset(s.strip() for s in raw.split(",") if s.strip())


def qa_from_status() -> str:
    return os.environ.get("DAILY_TASK_LOG_QA_FROM_STATUS", "Deployed on dev-int").strip()


def user_commented_on_issue(
    comments: list[dict[str, Any]],
    account_id: str,
) -> bool:
    want = (account_id or "").strip()
    if not want:
        return False
    for comment in comments:
        author_id = ((comment.get("author") or {}).get("accountId") or "").strip()
        if author_id == want:
            return True
    return False


def user_performed_qa_status_move(
    histories: list[dict[str, Any]],
    account_id: str,
    from_status: str,
    to_statuses: frozenset[str],
) -> bool:
    want = (account_id or "").strip()
    if not want or not from_status or not to_statuses:
        return False
    for history in histories:
        author_id = ((history.get("author") or {}).get("accountId") or "").strip()
        if author_id != want:
            continue
        for item in history.get("items") or []:
            if (item.get("field") or "").lower() != "status":
                continue
            if (item.get("fromString") or "").strip() != from_status:
                continue
            if (item.get("toString") or "").strip() in to_statuses:
                return True
    return False


def devops_qa_check_value(
    comments: list[dict[str, Any]],
    histories: list[dict[str, Any]],
    account_id: str,
) -> str:
    has_comment = user_commented_on_issue(comments, account_id)
    has_status_move = user_performed_qa_status_move(
        histories,
        account_id,
        qa_from_status(),
        qa_to_statuses(),
    )
    if has_comment or has_status_move:
        return ""
    return qa_missing_note()


def _header_fill() -> PatternFill:
    from openpyxl.styles.colors import Color

    return PatternFill(patternType="solid", fgColor=Color(theme=4, tint=0.6))


def _header_font() -> Font:
    return Font(bold=True, name="Calibri", size=11)


def weekday_lookback_days(today: date | None = None) -> int | None:
    """Return lookback days for today, or None on weekends."""
    today = today or date.today()
    if today.weekday() >= 5:
        return None
    return 3 if today.weekday() == 0 else 1


def quarter_sheet_name(for_date: date | None = None) -> str:
    """
    Calendar-quarter sheet tab from a date, e.g. Jul-Sept 26 (1 Jul–30 Sep 2026).
    Quarters: Jan–Mar, Apr–Jun, Jul–Sept, Oct–Dec. New tab on the 1st of each quarter.
    """
    for_date = for_date or date.today()
    year_suffix = str(for_date.year)[-2:]
    month = for_date.month
    if month <= 3:
        return f"Jan-Mar {year_suffix}"
    if month <= 6:
        return f"Apr-Jun {year_suffix}"
    if month <= 9:
        return f"Jul-Sept {year_suffix}"
    return f"Oct-Dec {year_suffix}"


def group_entries_by_quarter(entries: list[TaskEntry]) -> dict[str, list[TaskEntry]]:
    """Bucket rows by quarter sheet derived from each row's event date."""
    groups: dict[str, list[TaskEntry]] = {}
    for entry in entries:
        groups.setdefault(quarter_sheet_name(entry.event_date), []).append(entry)
    return groups


def format_dd_mmm(value: date | str) -> str:
    if isinstance(value, str):
        try:
            value = datetime.strptime(value[:10], "%Y-%m-%d").date()
        except ValueError:
            return value
    return value.strftime("%d-%b")


def jira_get_myself(base_url: str, email: str, api_token: str) -> dict[str, Any]:
    url = f"{base_url.rstrip('/')}/rest/api/3/myself"
    resp = requests.get(
        url,
        auth=(email, api_token),
        headers={"Accept": "application/json"},
        timeout=60,
    )
    if not resp.ok:
        print(f"JIRA myself error {resp.status_code}: {resp.text}", file=sys.stderr)
        sys.exit(1)
    return resp.json()


def jira_get_issue_changelog(
    base_url: str,
    email: str,
    api_token: str,
    issue_key: str,
    max_results: int = 100,
) -> list[dict[str, Any]]:
    url = f"{base_url.rstrip('/')}/rest/api/3/issue/{issue_key}/changelog"
    auth = (email, api_token)
    headers = {"Accept": "application/json"}
    histories: list[dict[str, Any]] = []
    start_at = 0
    while True:
        resp = requests.get(
            url,
            auth=auth,
            headers=headers,
            params={"startAt": start_at, "maxResults": max_results},
            timeout=60,
        )
        if not resp.ok:
            print(
                f"JIRA changelog error {resp.status_code} for {issue_key}: {resp.text}",
                file=sys.stderr,
            )
            sys.exit(1)
        data = resp.json()
        batch = data.get("values") or []
        histories.extend(batch)
        total = int(data.get("total") or len(histories))
        start_at += len(batch)
        if start_at >= total or not batch:
            break
    return histories


def parse_jira_timestamp(raw: str) -> datetime | None:
    raw = (raw or "").strip()
    if not raw:
        return None
    try:
        if raw.endswith("Z"):
            raw = f"{raw[:-1]}+00:00"
        return datetime.fromisoformat(raw)
    except ValueError:
        return None


def board_transition_dates(
    histories: list[dict[str, Any]],
    *,
    since: datetime,
    account_id: str | None,
    to_statuses: tuple[str, ...],
    from_statuses: tuple[str, ...] | None,
) -> list[date]:
    """Distinct local dates when a configured board status transition happened."""
    want_to = frozenset(to_statuses)
    want_from = frozenset(from_statuses) if from_statuses else None
    days: set[date] = set()
    for history in histories:
        created = parse_jira_timestamp(history.get("created") or "")
        if created is None:
            continue
        if created.astimezone() < since.astimezone():
            continue
        if account_id:
            author_id = ((history.get("author") or {}).get("accountId") or "").strip()
            if author_id != account_id:
                continue
        for item in history.get("items") or []:
            if (item.get("field") or "").lower() != "status":
                continue
            to_status = (item.get("toString") or "").strip()
            if to_status not in want_to:
                continue
            from_status = (item.get("fromString") or "").strip()
            if want_from is not None and from_status not in want_from:
                continue
            if from_status == to_status:
                continue
            days.add(created.astimezone().date())
    return sorted(days)


def board_transition_clauses(
    to_statuses: tuple[str, ...],
    from_statuses: tuple[str, ...] | None,
) -> list[str]:
    clauses: list[str] = []
    if from_statuses is None:
        for to_status in to_statuses:
            clauses.append(f'status changed to "{to_status}"')
        return clauses
    for from_status in from_statuses:
        for to_status in to_statuses:
            if from_status == to_status:
                continue
            clauses.append(f'status changed from "{from_status}" to "{to_status}"')
    return clauses


def board_transition_jql(
    days: int,
    account_id: str | None,
    projects: tuple[str, ...],
    to_statuses: tuple[str, ...],
    from_statuses: tuple[str, ...] | None,
) -> str:
    by_clause = ""
    if account_id:
        by_clause = " AND status changed by currentUser()"
    project_list = ", ".join(projects)
    transition_clauses = board_transition_clauses(to_statuses, from_statuses)
    if not transition_clauses:
        transition_expr = 'status changed to "UAT/Staging"'
    elif len(transition_clauses) == 1:
        transition_expr = transition_clauses[0]
    else:
        transition_expr = f"({' OR '.join(transition_clauses)})"
    return (
        f"project in ({project_list})"
        f" AND {transition_expr}"
        f"{by_clause}"
        f" during (-{days}d, now())"
        " ORDER BY updated DESC"
    )


def fetch_board_entries(
    base_url: str,
    email: str,
    token: str,
    *,
    lookback_days: int,
    account_id: str | None,
    since: datetime,
    projects: tuple[str, ...],
    to_statuses: tuple[str, ...],
    from_statuses: tuple[str, ...] | None,
) -> list[TaskEntry]:
    jql = board_transition_jql(
        lookback_days,
        account_id,
        projects,
        to_statuses,
        from_statuses,
    )
    transition_label = describe_board_transitions(to_statuses, from_statuses)
    print(f"Board JQL ({', '.join(projects)}; {transition_label}): {jql}")
    issues = dedupe_issues_by_key(jira_search_issues(base_url, email, token, jql=jql))
    entries: list[TaskEntry] = []
    for issue in issues:
        key = (issue.get("key") or "").strip()
        if not key:
            continue
        summary = ((issue.get("fields") or {}).get("summary") or "").strip()
        histories = jira_get_issue_changelog(base_url, email, token, key)
        transition_days = board_transition_dates(
            histories,
            since=since,
            account_id=account_id,
            to_statuses=to_statuses,
            from_statuses=from_statuses,
        )
        for day in transition_days:
            entries.append(
                TaskEntry(
                    event_date=day,
                    task_id=key,
                    description=summary,
                    remarks="",
                    qa_check="",
                    sort_rank=0,
                ),
            )
    return entries


def saathi_release_first_calendar_day(
    comments: list[dict[str, Any]],
    release_author: str,
    release_prefix: str,
) -> str | None:
    """Oldest calendar date (YYYY-MM-DD) with a matching Saathi release comment."""
    days = saathi_release_calendar_days_newest_first(
        comments,
        release_author,
        release_prefix,
    )
    if not days:
        return None
    return min(days)


def devops_remarks(
    comments: list[dict[str, Any]],
    release_author: str,
    release_prefix: str,
) -> str:
    first_day = saathi_release_first_calendar_day(
        comments,
        release_author,
        release_prefix,
    )
    if first_day:
        return f"Release Done\n{format_dd_mmm(first_day)}"
    return "Release Pending"


def issue_updated_date(issue: dict[str, Any]) -> date:
    raw = ((issue.get("fields") or {}).get("updated") or "").strip()
    parsed = parse_jira_timestamp(raw)
    if parsed is None:
        return date.today()
    return parsed.astimezone().date()


def fetch_devops_entries(
    base_url: str,
    email: str,
    token: str,
    *,
    jql: str,
    account_id: str,
    comment_lookback_start: datetime,
    release_author: str,
    release_prefix: str,
    automation_author: str,
    automation_prefix: str,
    filter_creators: frozenset[str],
    filter_projects: tuple[str, ...],
) -> list[TaskEntry]:
    print(f"DEVOPS JQL: {jql}")
    issues = dedupe_issues_by_key(jira_search_issues(base_url, email, token, jql=jql))
    entries: list[TaskEntry] = []
    skipped_by_filter = 0
    skipped_stale_deployed = 0
    for issue in issues:
        key = (issue.get("key") or "").strip()
        if not key:
            continue
        if not release_issue_passes_creator_link_filter(
            issue,
            filter_creators,
            filter_projects,
        ):
            skipped_by_filter += 1
            continue
        comments = jira_get_issue_comments(base_url, email, token, key)
        if skip_automation_only_in_updated_window(
            comments,
            comment_lookback_start,
            release_author,
            release_prefix,
            automation_author,
            automation_prefix,
        ):
            skipped_stale_deployed += 1
            continue
        histories = jira_get_issue_changelog(base_url, email, token, key)
        summary = ((issue.get("fields") or {}).get("summary") or "").strip()
        first_release_day = saathi_release_first_calendar_day(
            comments,
            release_author,
            release_prefix,
        )
        event_day = (
            datetime.strptime(first_release_day, "%Y-%m-%d").date()
            if first_release_day
            else issue_updated_date(issue)
        )
        qa_check = devops_qa_check_value(comments, histories, account_id)
        entries.append(
            TaskEntry(
                event_date=event_day,
                task_id=key,
                description=summary,
                remarks=devops_remarks(comments, release_author, release_prefix),
                qa_check=qa_check,
                sort_rank=1,
            ),
        )
    if skipped_by_filter:
        print(f"Skipped {skipped_by_filter} DEVOPS ticket(s) (creator linked-project filter).")
    if skipped_stale_deployed:
        print(
            f"Skipped {skipped_stale_deployed} DEVOPS ticket(s) "
            "(Automation-only comment in updated window).",
        )
    return entries


def sheet_base_year(sheet_name: str, today: date | None = None) -> int:
    today = today or date.today()
    match = re.search(r"(\d{2})$", sheet_name.strip())
    if match:
        return 2000 + int(match.group(1))
    return today.year


def parse_sheet_date(value: Any, base_year: int) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%d-%b", "%d-%b-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(text, fmt)
            if fmt == "%d-%b":
                return parsed.replace(year=base_year).date()
            return parsed.date()
        except ValueError:
            continue
    return None


def read_sheet_entries(ws: Worksheet, base_year: int) -> list[TaskEntry]:
    entries: list[TaskEntry] = []
    current_date: date | None = None
    for row_idx in range(2, ws.max_row + 1):
        date_raw = ws.cell(row_idx, 1).value
        task_id = ws.cell(row_idx, 2).value
        description = ws.cell(row_idx, 3).value
        remarks = ws.cell(row_idx, 4).value
        qa_check = ws.cell(row_idx, 5).value

        if date_raw:
            parsed = parse_sheet_date(date_raw, base_year)
            if parsed:
                current_date = parsed

        if not task_id:
            continue
        if current_date is None:
            continue

        task_key = str(task_id).strip()
        sort_rank = 1 if task_key.upper().startswith("DEVOPS-") else 0
        entries.append(
            TaskEntry(
                event_date=current_date,
                task_id=task_key,
                description=str(description or "").strip(),
                remarks=str(remarks or "").strip(),
                qa_check=str(qa_check or "").strip(),
                sort_rank=sort_rank,
            ),
        )
    return entries


def merge_task_entries(
    existing: list[TaskEntry],
    incoming: list[TaskEntry],
) -> tuple[list[TaskEntry], int]:
    by_id: dict[str, TaskEntry] = {}
    for entry in existing:
        by_id[entry.task_id.upper()] = entry
    added = 0
    for entry in incoming:
        key = entry.task_id.upper()
        if key in by_id:
            continue
        by_id[key] = entry
        added += 1
    merged = sorted(by_id.values(), key=lambda e: (e.event_date, e.sort_rank, e.task_id))
    return merged, added


def write_entries_to_sheet(ws: Worksheet, entries: list[TaskEntry]) -> None:
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    ordered = sorted(entries, key=lambda e: (e.event_date, e.sort_rank, e.task_id))
    row_idx = 1
    for group_date, group in _group_by_date(ordered):
        for index, entry in enumerate(group):
            row_idx += 1
            date_value = format_dd_mmm(entry.event_date) if index == 0 else None
            ws.cell(row_idx, 1, date_value)
            ws.cell(row_idx, 2, entry.task_id)
            ws.cell(row_idx, 3, entry.description)
            ws.cell(row_idx, 4, entry.remarks or None)
            ws.cell(row_idx, 5, entry.qa_check or None)
        row_idx = _write_blank_row(ws, row_idx)


def apply_sheet_headers(ws: Worksheet) -> None:
    for col, header in enumerate(SHEET_HEADERS, start=1):
        cell = ws.cell(1, col)
        cell.value = header
        cell.font = _header_font()
        cell.fill = _header_fill()
    ws.column_dimensions["A"].width = 12.72
    ws.column_dimensions["B"].width = 15.85
    ws.column_dimensions["C"].width = 52.15
    ws.column_dimensions["D"].width = 25.28
    ws.column_dimensions["E"].width = 22.0


def append_entries_to_sheet(
    ws: Worksheet,
    entries: list[TaskEntry],
    *,
    sheet_name: str,
) -> int:
    if not entries:
        return 0

    base_year = sheet_base_year(sheet_name)
    existing = read_sheet_entries(ws, base_year)
    merged, added = merge_task_entries(existing, entries)
    if added == 0:
        print("All candidate rows already exist in the sheet (by TaskId).")
        return 0

    write_entries_to_sheet(ws, merged)
    return added


def ensure_sheet(wb: Any, sheet_name: str) -> Worksheet:
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        apply_sheet_headers(ws)
        return ws
    ws = wb.create_sheet(sheet_name)
    apply_sheet_headers(ws)
    print(f"Created sheet tab: {sheet_name}")
    return ws


def _group_by_date(entries: list[TaskEntry]) -> list[tuple[date, list[TaskEntry]]]:
    if not entries:
        return []
    groups: list[tuple[date, list[TaskEntry]]] = []
    current_date = entries[0].event_date
    bucket: list[TaskEntry] = []
    for entry in entries:
        if entry.event_date != current_date:
            groups.append((current_date, bucket))
            current_date = entry.event_date
            bucket = []
        bucket.append(entry)
    groups.append((current_date, bucket))
    return groups


def _write_blank_row(ws: Worksheet, row_idx: int) -> int:
    row_idx += 1
    for col in range(1, SHEET_COLUMN_COUNT + 1):
        ws.cell(row_idx, col, None)
    return row_idx


def append_to_workbook(path: Path, sheet_name: str, entries: list[TaskEntry]) -> int:
    if path.is_file():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        default = wb.active
        if default is not None:
            wb.remove(default)
    ws = ensure_sheet(wb, sheet_name)
    wrote = append_entries_to_sheet(ws, entries, sheet_name=sheet_name)
    if wrote:
        wb.save(path)
    return wrote


def build_since_datetime(lookback_days: int) -> datetime:
    start_day = date.today() - timedelta(days=lookback_days)
    local_tz = datetime.now().astimezone().tzinfo
    return datetime.combine(start_day, datetime.min.time(), tzinfo=local_tz)


def main() -> None:
    load_dotenv()
    parser = argparse.ArgumentParser(
        description="Append daily COR/LINUX/ALPINE and DEVOPS release tasks to DailyTaskLogs.xlsx.",
    )
    parser.add_argument(
        "--days",
        type=int,
        metavar="N",
        help="Override lookback days (default: 3 on Monday, 1 Tue–Fri).",
    )
    parser.add_argument(
        "--workbook",
        default=os.environ.get("DAILY_TASK_LOG_WORKBOOK", str(DEFAULT_WORKBOOK)),
        help="Path to DailyTaskLogs.xlsx",
    )
    parser.add_argument(
        "--sheet",
        default="",
        help="Worksheet tab (default: current quarter, e.g. Jul-Sept 26).",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Print rows without writing the workbook.",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Run even on Saturday/Sunday.",
    )
    parser.add_argument(
        "--reorder",
        action="store_true",
        help="Rewrite the quarter sheet in date order from existing rows only.",
    )
    args = parser.parse_args()

    sheet_override = args.sheet.strip()
    default_sheet = sheet_override or quarter_sheet_name()
    workbook_path = Path(args.workbook)

    if args.reorder:
        if not workbook_path.is_file():
            print(f"Workbook not found: {workbook_path}", file=sys.stderr)
            sys.exit(1)
        wb = load_workbook(workbook_path)
        ws = ensure_sheet(wb, default_sheet)
        base_year = sheet_base_year(default_sheet)
        existing = read_sheet_entries(ws, base_year)
        merged, _ = merge_task_entries(existing, [])
        write_entries_to_sheet(ws, merged)
        wb.save(workbook_path)
        print(f"Reordered {len(merged)} row(s) on {default_sheet}.")
        return

    lookback_days = args.days
    if lookback_days is None:
        lookback_days = weekday_lookback_days()
        if lookback_days is None and not args.force:
            print("Weekend — daily task log runs Monday through Friday only.")
            return
        if lookback_days is None:
            lookback_days = 1

    base_url = _require_env("JIRA_BASE_URL")
    email = _require_env("JIRA_EMAIL")
    token = _require_env("JIRA_API_TOKEN")

    only_my_transitions = os.environ.get(
        "DAILY_TASK_LOG_ONLY_MY_TRANSITIONS",
        "1",
    ).strip().lower() in ("1", "true", "yes")
    account_id = (jira_get_myself(base_url, email, token).get("accountId") or "").strip()
    if not account_id:
        print("Could not resolve JIRA accountId for QA checks.", file=sys.stderr)
        sys.exit(1)
    if only_my_transitions:
        print(f"Board transitions filtered to current user ({email}).")
    else:
        print(f"DEVOPS QA checks use current user ({email}); board transitions include all users.")

    board_account_filter = account_id if only_my_transitions else None
    board_project_list = board_projects()
    board_to_list = board_to_statuses()
    board_from_list = board_from_statuses(board_to_list)

    since = build_since_datetime(lookback_days)
    print(f"Lookback: {lookback_days} day(s) since {since.date()}")
    print(f"Board projects: {', '.join(board_project_list)}")
    print(
        "Board transitions: "
        f"{describe_board_transitions(board_to_list, board_from_list)}"
    )

    board_entries = fetch_board_entries(
        base_url,
        email,
        token,
        lookback_days=lookback_days,
        account_id=board_account_filter,
        since=since,
        projects=board_project_list,
        to_statuses=board_to_list,
        from_statuses=board_from_list,
    )

    default_devops_jql = (
        f'project = DEVOPS AND updated >= -{lookback_days}d '
        'AND issuetype = Release ORDER BY updated DESC'
    )
    devops_jql = os.environ.get("DAILY_TASK_LOG_JQL", "").strip() or default_devops_jql
    devops_jql, lookback_field = apply_jql_lookback_days(devops_jql, lookback_days)
    print(f"Using DAILY_TASK_LOG_JQL with {lookback_field} >= -{lookback_days}d")

    env_lookback = os.environ.get("JIRA_COMMENT_LOOKBACK_DAYS", "").strip()
    default_comment_lookback = int(env_lookback) if env_lookback.isdigit() else lookback_days
    comment_lookback_start = jql_updated_lookback_start(
        devops_jql,
        default_days=default_comment_lookback,
    )

    release_author = os.environ.get("JIRA_RELEASE_COMMENT_AUTHOR", "Saathi").strip()
    release_prefix = os.environ.get(
        "JIRA_RELEASE_COMMENT_PREFIX",
        "Release has been completed  for",
    ).strip()
    automation_author = (
        os.environ.get("JIRA_DEPLOYED_AUTOMATION_COMMENT_AUTHOR", "").strip()
        or "Automation for Jira"
    )
    automation_prefix = (
        os.environ.get("JIRA_DEPLOYED_AUTOMATION_COMMENT_PREFIX", "").strip()
        or "Ticket Automation Executed. Your ticket has been marked as"
    )
    filter_creators = frozenset(parse_env_csv("JIRA_RELEASE_LINK_FILTER_CREATORS"))
    filter_projects = tuple(parse_env_csv("JIRA_RELEASE_LINK_FILTER_PROJECTS"))

    devops_entries = fetch_devops_entries(
        base_url,
        email,
        token,
        jql=devops_jql,
        account_id=account_id,
        comment_lookback_start=comment_lookback_start,
        release_author=release_author,
        release_prefix=release_prefix,
        automation_author=automation_author,
        automation_prefix=automation_prefix,
        filter_creators=filter_creators,
        filter_projects=filter_projects,
    )

    entries = board_entries + devops_entries
    if not entries:
        print("No new tasks found for the current lookback window.")
        return

    entry_groups = (
        {sheet_override: entries}
        if sheet_override
        else group_entries_by_quarter(entries)
    )

    print(f"Workbook: {args.workbook}")
    if sheet_override:
        print(f"Sheet tab: {sheet_override}")
    else:
        quarter_tabs = ", ".join(sorted(entry_groups))
        print(f"Sheet tab(s) (auto quarterly): {quarter_tabs}")
    print(f"Found {len(board_entries)} board transition(s) and {len(devops_entries)} DEVOPS ticket(s).")

    if args.dry_run:
        for entry in sorted(entries, key=lambda e: (e.event_date, e.sort_rank, e.task_id)):
            qa_preview = entry.qa_check.replace("\n", " / ") if entry.qa_check else ""
            target_sheet = sheet_override or quarter_sheet_name(entry.event_date)
            print(
                f"{format_dd_mmm(entry.event_date):>7}  {entry.task_id:<14}  "
                f"[{target_sheet}]  "
                f"{entry.description[:50]}  {entry.remarks!r}  {qa_preview!r}",
            )
        return

    total_wrote = 0
    for sheet_name in sorted(entry_groups):
        sheet_entries = entry_groups[sheet_name]
        wrote = append_to_workbook(workbook_path, sheet_name, sheet_entries)
        total_wrote += wrote
        if wrote:
            print(f"Appended {wrote} row(s) to {args.workbook} → {sheet_name}.")

    if total_wrote == 0:
        print("No new rows written (duplicates or empty after filters).")


if __name__ == "__main__":
    main()
